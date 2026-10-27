"""Document intake (FR-01) and content preparation (FR-02).

Validates file type, size and integrity, then extracts text into
page/section-aware units with document metadata (Doc-ID, title,
category, version, status, effective date).

Synthetic demo documents declare metadata in a header block::

    Doc-ID: HR-POL-007
    Title: Data Retention Policy
    Category: HR
    Version: 2.0
    Status: effective
    Effective: 2025-01-01
"""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field, replace
from itertools import groupby
from pathlib import Path
from typing import Any

from src.core.exceptions import ValidationError
from src.core.logging import get_logger
from src.knowledge.vision import (
    VisionClient,
    VisionResult,
    load_prompt,
    parse_json_reply,
    to_png,
)

logger = get_logger(__name__)

ALLOWED_EXTENSIONS: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}
IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg"})

# File signatures ("magic bytes"). The content must match the extension.
_MAGIC: dict[str, tuple[bytes, ...]] = {
    ".pdf": (b"%PDF",),
    ".docx": (b"PK\x03\x04",),
    ".png": (b"\x89PNG\r\n\x1a\n",),
    ".jpg": (b"\xff\xd8\xff",),
    ".jpeg": (b"\xff\xd8\xff",),
}
_EXECUTABLE_MAGIC = (b"MZ", b"\x7fELF", b"\xcf\xfa\xed\xfe", b"\xfe\xed\xfa")
_KIND_NAMES = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".png": "PNG image",
    ".jpg": "JPEG image",
}
VALID_STATUSES = ("effective", "superseded", "draft")
_DOC_ID_RE = re.compile(r"[A-Za-z]{2,6}(?:-[A-Za-z0-9]+)+")

_META_RE = re.compile(
    r"^\s*(doc-id|title|category|version|status|effective|owner)\s*:\s*(.+?)\s*$",
    re.IGNORECASE,
)
_NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+[A-Z][^.!?]{2,80}$")
_MIN_TEXT_CHARS = 20


@dataclass
class TextUnit:
    """A contiguous piece of text with its location in the source.

    content_type: text | table | chart | ocr | image
    source_method: text_layer | vision_ocr | docx | plain
    """

    text: str
    page: int | None = None
    section: str | None = None
    content_type: str = "text"
    source_method: str = "plain"
    asset_id: str | None = None
    bbox: list[float] | None = None


@dataclass
class Asset:
    """An image kept so the UI can show the original (page render, chart)."""

    asset_id: str
    kind: str  # page | image | chart
    page: int | None
    png: bytes
    width: int
    height: int


@dataclass
class ParsedDocument:
    """Validated, extracted document ready for chunking."""

    filename: str
    mime_type: str
    size_bytes: int
    metadata: dict[str, str] = field(default_factory=dict)
    units: list[TextUnit] = field(default_factory=list)
    pages: int | None = None
    assets: list[Asset] = field(default_factory=list)
    vision_input_tokens: int = 0
    vision_output_tokens: int = 0
    vision_calls: int = 0
    skipped_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# ── Validation ───────────────────────────────────────────────


def validate_upload(filename: str, content: bytes, max_bytes: int) -> str:
    """Validate an upload and return its canonical MIME type.

    Raises:
        ValidationError: unsupported type, empty, too large or corrupted.
    """
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValidationError(
            f"Unsupported file type '{ext or 'none'}'. "
            f"Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )
    if not content:
        raise ValidationError("File is empty.")
    if len(content) > max_bytes:
        raise ValidationError(file_too_large_message(len(content), max_bytes))
    if content.startswith(_EXECUTABLE_MAGIC):
        raise ValidationError(
            f"File has a {ext} extension but is an executable program. "
            "Renamed executables are rejected."
        )
    signatures = _MAGIC.get(ext)
    if signatures and not content.startswith(signatures):
        kind = _KIND_NAMES.get(ext, ext)
        raise ValidationError(
            f"File has a {ext} extension but is not a valid {kind} "
            "(content does not match the file type, or it is corrupted)."
        )
    if ext in (".txt", ".md") and b"\x00" in content:
        raise ValidationError("Text file contains binary data (corrupted?).")
    return ALLOWED_EXTENSIONS[ext]


def file_too_large_message(size: int, max_bytes: int) -> str:
    return (
        f"File is too large ({size / 1_048_576:.1f} MB). "
        f"Limit is {max_bytes / 1_048_576:.0f} MB."
    )


def normalize_metadata(overrides: dict[str, str | None] | None) -> dict[str, str]:
    """Validate metadata supplied with an upload (form fields)."""
    clean: dict[str, str] = {}
    for key, value in (overrides or {}).items():
        if value is None or not str(value).strip():
            continue
        value = str(value).strip()[:100]
        if key == "status":
            value = value.lower()
            if value not in VALID_STATUSES:
                raise ValidationError(
                    f"Status must be one of: {', '.join(VALID_STATUSES)}."
                )
        if key == "effective" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValidationError("Effective date must look like YYYY-MM-DD.")
        clean[key] = value
    return clean


def doc_id_from_filename(filename: str) -> str:
    """'TRV-POL-003_Travel.pdf' -> 'TRV-POL-003'; 'remote_work.txt' -> 'REMOTE-WORK'."""
    stem = Path(filename).stem
    first = stem.split("_")[0]
    if _DOC_ID_RE.fullmatch(first):
        return first.upper()
    return re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-").upper()[:60] or "DOC"


# ── Extraction ───────────────────────────────────────────────

_OCR_MIN_PAGE_CHARS = 30  # below this a PDF page is treated as scanned
_IMAGE_MIN_PX = 150  # smaller embedded images (icons, logos) are ignored
_VECTOR_CHART_DRAWINGS = 200  # many vector paths + little text = drawn chart
_VECTOR_CHART_MAX_TEXT = 300
_RENDER_DPI = 150


class _VisionBudget:
    """Caps vision calls per document and tallies their token usage."""

    def __init__(
        self, vision: VisionClient | None, max_calls: int, doc: ParsedDocument
    ):
        self.vision = vision
        self.max_calls = max_calls
        self.doc = doc

    @property
    def available(self) -> bool:
        return self.vision is not None and self.vision.enabled

    def call(self, image: bytes, prompt: str, page: int) -> VisionResult | None:
        """Return the result, or None when the page was skipped (cap or no vision)."""
        if not self.available:
            if page not in self.doc.skipped_pages:
                self.doc.skipped_pages.append(page)
            return None
        if self.doc.vision_calls >= self.max_calls:
            if page not in self.doc.skipped_pages:
                self.doc.skipped_pages.append(page)
            _warn_once(
                self.doc,
                f"Vision call limit reached ({self.max_calls} per document); "
                "remaining scanned pages and charts were skipped.",
            )
            return None
        assert self.vision is not None
        result = self.vision.describe(image, prompt)
        self.doc.vision_calls += 1
        self.doc.vision_input_tokens += result.input_tokens
        self.doc.vision_output_tokens += result.output_tokens
        if not result.ok:
            _warn_once(self.doc, f"Vision failed on page {page}: {result.error}")
        return result


def _warn_once(doc: ParsedDocument, message: str) -> None:
    if message not in doc.warnings:
        doc.warnings.append(message)


def parse_document(
    filename: str,
    content: bytes,
    max_bytes: int,
    overrides: dict[str, str | None] | None = None,
    vision: VisionClient | None = None,
    vision_max_calls: int = 30,
) -> ParsedDocument:
    """Validate and extract a document into located text units.

    ``overrides`` (category, version, status, effective) come from the upload
    form and win over the file's own header; the filename is the last fallback.
    ``vision`` reads scanned pages, image uploads and charts; without it those
    are skipped and reported in ``skipped_pages`` / ``warnings``.
    """
    mime = validate_upload(filename, content, max_bytes)
    ext = Path(filename).suffix.lower()
    form_meta = normalize_metadata(overrides)
    doc = ParsedDocument(filename=filename, mime_type=mime, size_bytes=len(content))
    budget = _VisionBudget(vision, vision_max_calls, doc)

    if ext == ".pdf":
        units, doc.pages = _extract_pdf(content, budget)
    elif ext == ".docx":
        units = _extract_docx(content)
    elif ext in IMAGE_EXTENSIONS:
        units, doc.pages = _extract_image(content, budget), 1
    else:
        text = _decode_text(content)
        units = _split_tables(
            _extract_text(text, markdown=ext == ".md"), source_method="plain"
        )

    metadata, units = _pull_metadata(units)
    units = [u for u in units if u.text.strip()]

    total_chars = sum(len(u.text.strip()) for u in units)
    if total_chars < _MIN_TEXT_CHARS:
        if ext == ".pdf" and not budget.available:
            raise ValidationError(
                "PDF contains no extractable text. Scanned or handwritten pages "
                f"need vision OCR, and {VisionClient.UNAVAILABLE}."
            )
        if ext == ".pdf":
            raise ValidationError(
                "PDF contains no extractable text, and OCR returned no text."
            )
        raise ValidationError("Document contains no meaningful text.")

    stem = Path(filename).stem
    metadata.update(form_meta)
    if "status" in metadata:
        metadata["status"] = metadata["status"].lower()
    metadata.setdefault("doc-id", doc_id_from_filename(filename))
    metadata.setdefault("title", stem.replace("_", " "))
    doc.metadata = metadata
    doc.units = units
    return doc


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def _heading_of(line: str, markdown: bool) -> str | None:
    """Return the heading text if the line looks like a section heading."""
    stripped = line.strip()
    if not stripped:
        return None
    if markdown and stripped.startswith("#"):
        return stripped.lstrip("#").strip() or None
    if _NUMBERED_HEADING_RE.match(stripped):
        return stripped
    if (
        stripped.isupper()
        and 3 <= len(stripped) <= 60
        and any(c.isalpha() for c in stripped)
        # "TOTAL USD 367.20" is content, not a heading.
        and not any(c.isdigit() for c in stripped)
        and not _META_RE.match(stripped)
    ):
        return stripped.title()
    return None


def _split_sections(
    text: str, page: int | None, current: str | None, markdown: bool
) -> tuple[list[TextUnit], str | None]:
    """Split text at heading lines, carrying the active section forward."""
    units: list[TextUnit] = []
    buf: list[str] = []
    for line in text.splitlines():
        heading = _heading_of(line, markdown)
        if heading:
            if "".join(buf).strip():
                units.append(TextUnit("\n".join(buf).strip(), page, current))
            buf = []
            current = heading
        else:
            buf.append(line)
    if "".join(buf).strip():
        units.append(TextUnit("\n".join(buf).strip(), page, current))
    return units, current


def _extract_text(text: str, markdown: bool) -> list[TextUnit]:
    units, _ = _split_sections(text, None, None, markdown)
    return units


def _is_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _split_tables(units: list[TextUnit], source_method: str) -> list[TextUnit]:
    """Pull markdown pipe tables out of text units as their own table units."""
    out: list[TextUnit] = []
    for unit in units:
        unit.source_method = source_method
        out.extend(_split_unit_tables(unit))
    return out


def _split_unit_tables(unit: TextUnit) -> list[TextUnit]:
    blocks: list[tuple[bool, list[str]]] = []
    for is_table, group in groupby(unit.text.splitlines(), key=_is_table_line):
        lines = list(group)
        is_table = is_table and len(lines) >= 2  # one pipe line is just text
        if blocks and not is_table and not blocks[-1][0]:
            blocks[-1][1].extend(lines)
        else:
            blocks.append((is_table, lines))
    out = []
    for is_table, lines in blocks:
        text = "\n".join(ln.strip() if is_table else ln for ln in lines).strip()
        if text:
            kind = "table" if is_table else unit.content_type
            out.append(replace(unit, text=text, content_type=kind))
    return out


def rows_to_markdown(rows: list[list[str]]) -> str:
    """First row is the header. Pipes and newlines inside cells are escaped."""
    clean = [
        [" ".join((c or "").split()).replace("|", "/") for c in row] for row in rows
    ]
    clean = [r for r in clean if any(r)]
    if not clean:
        return ""
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return "\n".join(lines)


def split_table(markdown: str, max_chars: int) -> list[str]:
    """Split a long markdown table by rows, repeating the header in each piece."""
    lines = markdown.strip().splitlines()
    if len(markdown) <= max_chars or len(lines) <= 3:
        return [markdown.strip()]
    has_sep = len(lines) > 1 and set(lines[1].replace(" ", "")) <= set("|-:")
    header = lines[:2] if has_sep else lines[:1]
    pieces: list[list[str]] = []
    current: list[str] = []
    for row in lines[len(header) :]:
        if current and len("\n".join(header + current + [row])) > max_chars:
            pieces.append(header + current)
            current = []
        current.append(row)
    if current:
        pieces.append(header + current)
    return ["\n".join(p) for p in pieces]


# ── PDF (PyMuPDF): text layer, tables, scanned pages, charts ──


def _extract_pdf(content: bytes, budget: _VisionBudget) -> tuple[list[TextUnit], int]:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise ValidationError(f"PDF could not be parsed (corrupted?): {exc}") from exc
    if pdf.needs_pass:
        raise ValidationError("PDF is password-protected and cannot be read.")

    units: list[TextUnit] = []
    section: str | None = None
    try:
        for page_no, page in enumerate(pdf, start=1):
            raw_text = page.get_text().strip()
            if len(raw_text) < _OCR_MIN_PAGE_CHARS:
                # Scanned or handwritten page: the whole page goes to vision OCR.
                units.extend(_ocr_page(page, page_no, section, budget))
                continue
            page_units, section = _page_text_and_tables(page, page_no, section)
            units.extend(page_units)
            units.extend(_read_images(page, page_no, section, budget))
            drawings = len(page.get_drawings())
            if (
                drawings > _VECTOR_CHART_DRAWINGS
                and len(raw_text) < _VECTOR_CHART_MAX_TEXT
            ):
                logger.info(
                    "kis_vector_chart_heuristic", page=page_no, drawings=drawings
                )
                png, w, h = _render(page)
                units.extend(_chart_units(png, w, h, page_no, section, budget, None))
        return units, pdf.page_count
    finally:
        pdf.close()


def _page_text_and_tables(
    page: Any, page_no: int, section: str | None
) -> tuple[list[TextUnit], str | None]:
    """Text blocks and tables in top-to-bottom order.

    Processing in reading order means a table is filed under the heading
    above it, and text inside a table is not indexed twice.
    """
    import pymupdf

    try:
        tables = page.find_tables().tables
    except Exception:
        tables = []
    rects = [pymupdf.Rect(t.bbox) for t in tables]
    items: list[tuple[float, str, Any]] = [
        (b[1], "text", b[4])
        for b in page.get_text("blocks", sort=True)
        if b[6] == 0 and not any(pymupdf.Rect(b[:4]).intersects(r) for r in rects)
    ]
    items += [(t.bbox[1], "table", t) for t in tables]
    items.sort(key=lambda item: item[0])

    units: list[TextUnit] = []
    text_buf: list[str] = []
    for _, kind, payload in items:
        if kind == "text":
            text_buf.append(payload)
            continue
        found, section = _split_sections("\n".join(text_buf), page_no, section, False)
        units.extend(found)
        text_buf = []
        md = rows_to_markdown(_table_cells(page, payload))
        if md:
            units.append(
                TextUnit(
                    md,
                    page_no,
                    section,
                    content_type="table",
                    bbox=[round(v, 1) for v in payload.bbox],
                )
            )
    found, section = _split_sections("\n".join(text_buf), page_no, section, False)
    units.extend(found)
    for unit in units:
        unit.source_method = "text_layer"
    return units, section


def _table_cells(page: Any, table: Any) -> list[list[str]]:
    """Cell text read from the page by clip (keeps '120,000' intact)."""
    import pymupdf

    try:
        return [
            [
                page.get_text("text", clip=pymupdf.Rect(cell)).strip() if cell else ""
                for cell in row.cells
            ]
            for row in table.rows
        ]
    except Exception:
        return table.extract()


def _render(page: Any, clip: Any = None) -> tuple[bytes, int, int]:
    pix = page.get_pixmap(dpi=_RENDER_DPI, clip=clip)
    return pix.tobytes("png"), pix.width, pix.height


def _add_asset(
    budget: _VisionBudget, kind: str, page: int | None, png: bytes, w: int, h: int
) -> str:
    asset_id = str(uuid.uuid4())
    budget.doc.assets.append(Asset(asset_id, kind, page, png, w, h))
    return asset_id


def _ocr_units(
    text: str, page: int | None, section: str | None, asset_id: str
) -> list[TextUnit]:
    units, _ = _split_sections(text, page, section, markdown=True)
    units = _split_tables(units, source_method="vision_ocr")
    for unit in units:
        if unit.content_type == "text":
            unit.content_type = "ocr"
        unit.asset_id = asset_id
    return units


def _ocr_page(
    page: Any, page_no: int, section: str | None, budget: _VisionBudget
) -> list[TextUnit]:
    png, w, h = _render(page)
    result = budget.call(png, load_prompt("ocr_v1"), page_no)
    if result is None or not result.ok or not result.text.strip():
        if result is not None and result.ok:
            _warn_once(budget.doc, f"OCR found no text on page {page_no}.")
        if page_no not in budget.doc.skipped_pages:
            budget.doc.skipped_pages.append(page_no)
        if not budget.available:
            _warn_once(
                budget.doc,
                f"Pages without a text layer were skipped: {VisionClient.UNAVAILABLE}.",
            )
        return []
    asset_id = _add_asset(budget, "page", page_no, png, w, h)
    return _ocr_units(result.text, page_no, section, asset_id)


def _read_images(
    page: Any, page_no: int, section: str | None, budget: _VisionBudget
) -> list[TextUnit]:
    """Charts and photos embedded as images on a text page."""
    units: list[TextUnit] = []
    for info in page.get_images(full=True):
        xref, width, height = info[0], info[2], info[3]
        if width < _IMAGE_MIN_PX or height < _IMAGE_MIN_PX:
            continue
        rects = page.get_image_rects(xref)
        if not rects:
            continue
        if not budget.available:
            _warn_once(
                budget.doc,
                f"Charts and images were not read: {VisionClient.UNAVAILABLE}.",
            )
            return units
        png, w, h = _render(page, clip=rects[0])
        bbox = [round(v, 1) for v in rects[0]]
        units.extend(_chart_units(png, w, h, page_no, section, budget, bbox))
    return units


def chart_markdown(series: list[dict[str, Any]]) -> str:
    """Chart series -> markdown table (x label, one column per series)."""
    labels: list[str] = []
    values: dict[str, dict[str, Any]] = {}
    for s in series:
        name = str(s.get("name") or "value")
        for point in s.get("points") or []:
            if not isinstance(point, list | tuple) or len(point) < 2:
                continue
            x = str(point[0])
            if x not in labels:
                labels.append(x)
            values.setdefault(name, {})[x] = point[1]
    names = list(values)
    if not names:
        return ""
    rows = [["label", *names]] + [
        [x, *[str(values[n].get(x, "")) for n in names]] for x in labels
    ]
    return rows_to_markdown(rows)


def _chart_units(
    png: bytes,
    w: int,
    h: int,
    page_no: int,
    section: str | None,
    budget: _VisionBudget,
    bbox: list[float] | None,
) -> list[TextUnit]:
    result = budget.call(png, load_prompt("chart_v1"), page_no)
    if result is None or not result.ok:
        return []
    data = parse_json_reply(result.text) or {}
    kind = data.get("kind")
    if kind == "chart":
        asset_id = _add_asset(budget, "chart", page_no, png, w, h)
        title = str(data.get("title") or "Untitled chart")
        table = chart_markdown(data.get("series") or [])
        axes = f"x: {data.get('x_axis') or '-'}; y: {data.get('y_axis') or '-'}"
        text = (
            f"Chart: {title} (page {page_no})\n"
            f"Type: {data.get('chart_type') or 'chart'}. Axes: {axes}.\n"
            f"{str(data.get('summary') or '').strip()}\n"
            f"Data:\n{table}"
        )
        return [
            TextUnit(
                text.strip(),
                page_no,
                section,
                content_type="chart",
                source_method="vision_ocr",
                asset_id=asset_id,
                bbox=bbox,
            )
        ]
    if kind == "photo" and data.get("caption"):
        asset_id = _add_asset(budget, "image", page_no, png, w, h)
        return [
            TextUnit(
                f"Image (page {page_no}): {data['caption']}",
                page_no,
                section,
                content_type="image",
                source_method="vision_ocr",
                asset_id=asset_id,
                bbox=bbox,
            )
        ]
    return []  # logos, diagrams and other images are not indexed


# ── Images (photos of handwriting, scans) ────────────────────


def _extract_image(content: bytes, budget: _VisionBudget) -> list[TextUnit]:
    try:
        png, w, h = to_png(content)
    except Exception as exc:
        raise ValidationError(f"Image could not be read (corrupted?): {exc}") from exc
    if not budget.available:
        raise ValidationError(
            f"Images need vision OCR, and {VisionClient.UNAVAILABLE}."
        )
    result = budget.call(content, load_prompt("ocr_v1"), 1)
    if result is None or not result.ok:
        reason = result.error if result else "vision call limit reached"
        raise ValidationError(f"Could not read the image: {reason}")
    if not result.text.strip():
        raise ValidationError("OCR found no text in the image.")
    asset_id = _add_asset(budget, "image", 1, png, w, h)
    return _ocr_units(result.text, 1, None, asset_id)


# ── DOCX: paragraphs and tables in document order ────────────


def _extract_docx(content: bytes) -> list[TextUnit]:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ValidationError(f"DOCX could not be parsed (corrupted?): {exc}") from exc

    units: list[TextUnit] = []
    section: str | None = None
    buf: list[str] = []

    def flush() -> None:
        if buf:
            units.append(TextUnit("\n".join(buf), None, section, source_method="docx"))
        buf.clear()

    # Walk the body in order so every table stays under its own heading.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            style = (para.style.name if para.style is not None else "") or ""
            text = para.text.strip()
            if not text:
                continue
            if style.lower().startswith(("heading", "title")):
                flush()
                section = text
            else:
                buf.append(text)
        elif tag == "tbl":
            flush()
            table = Table(child, document)
            md = rows_to_markdown([[c.text for c in row.cells] for row in table.rows])
            if md:
                units.append(
                    TextUnit(
                        md, None, section, content_type="table", source_method="docx"
                    )
                )
    flush()
    return units


def _pull_metadata(units: list[TextUnit]) -> tuple[dict[str, str], list[TextUnit]]:
    """Extract ``Key: value`` metadata lines from the start of the document."""
    metadata: dict[str, str] = {}
    cleaned: list[TextUnit] = []
    for idx, unit in enumerate(units):
        if idx > 1 or unit.content_type != "text":
            cleaned.append(unit)
            continue
        kept: list[str] = []
        for line in unit.text.splitlines():
            match = _META_RE.match(line)
            if match and match.group(1).lower() not in metadata:
                metadata[match.group(1).lower()] = match.group(2)
            else:
                kept.append(line)
        cleaned.append(replace(unit, text="\n".join(kept).strip()))
    if "status" in metadata:
        metadata["status"] = metadata["status"].lower()
    return metadata, cleaned
